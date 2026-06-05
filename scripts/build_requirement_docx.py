from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT_PATH = Path("RAG_API_Service_Assessment_Requirement.docx")


BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
MUTED = RGBColor(89, 89, 89)
LIGHT_GRAY = "F2F4F7"
BORDER = "D9E2F3"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)

    for margin_name, value in {
        "top": top,
        "start": start,
        "bottom": bottom,
        "end": end,
    }.items():
        node = tc_mar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)

    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        border = borders.find(qn(tag))
        if border is None:
            border = OxmlElement(tag)
            borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "6")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), BORDER)


def set_table_width(table, width_dxa: int = 9360, indent_dxa: int = 120) -> None:
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(width_dxa))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")


def format_table(table, widths: list[int]) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    set_table_width(table)
    set_table_borders(table)

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row_index, row in enumerate(table.rows):
        for cell_index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_width(cell, widths[cell_index])
            set_cell_margins(cell)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.1
                for run in paragraph.runs:
                    run.font.name = "Calibri"
                    run.font.size = Pt(10)
            if row_index == 0:
                set_cell_shading(cell, LIGHT_GRAY)
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True
                        run.font.color.rgb = DARK_BLUE


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[int]):
    table = doc.add_table(rows=1, cols=len(headers))
    header_cells = table.rows[0].cells
    for index, header in enumerate(headers):
        header_cells[index].text = header
    for row_values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row_values):
            cells[index].text = value
    format_table(table, widths)
    return table


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.add_run(item)


def add_numbers(doc: Document, items: list[str]) -> None:
    for item in items:
        paragraph = doc.add_paragraph(style="List Number")
        paragraph.add_run(item)


def add_h1(doc: Document, text: str) -> None:
    doc.add_paragraph(text, style="Heading 1")


def add_h2(doc: Document, text: str) -> None:
    doc.add_paragraph(text, style="Heading 2")


def build_doc() -> None:
    doc = Document()
    section = doc.sections[0]
    section.start_type = WD_SECTION.NEW_PAGE
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    for style_name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for list_style_name in ("List Bullet", "List Number"):
        style = styles[list_style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.167

    header = section.header.paragraphs[0]
    header.text = "Local Extractive RAG API Service | Requirement Assessment"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.name = "Calibri"
        run.font.size = Pt(9)
        run.font.color.rgb = MUTED

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(3)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = title.add_run("Local Extractive RAG API Service")
    run.font.name = "Calibri"
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = RGBColor(11, 37, 69)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    run = subtitle.add_run("Requirement Document for Codex Implementation and Assessment")
    run.font.name = "Calibri"
    run.font.size = Pt(12)
    run.font.color.rgb = MUTED

    add_table(
        doc,
        ["Item", "Requirement"],
        [
            ["Project type", "Small Python backend API service"],
            ["Main goal", "Answer questions from a provided local plain-text document set using RAG principles"],
            ["Required endpoints", "POST /index and POST /ask"],
            ["Major constraint", "No external LLM APIs or external paid services"],
            ["Assessment level", "L1-level exercise; correctness, judgment, and readability matter most"],
        ],
        [2600, 6760],
    )

    add_h1(doc, "1. Objective")
    doc.add_paragraph(
        "Build a small API service that answers questions from a provided folder of plain-text documents "
        "about a fictional product and its policies. The service must demonstrate retrieval-augmented "
        "generation principles while staying fully local and simple."
    )
    doc.add_paragraph(
        "The implementation should be judged according to this requirement document. If a behavior is not "
        "supported by evidence from the indexed documents, the service must refuse to answer instead of "
        "inventing content."
    )

    add_h1(doc, "2. Scope")
    add_bullets(
        doc,
        [
            "Use Python for the backend service.",
            "Use a local docs/ folder containing .txt files as the document source.",
            "Build an in-memory retrieval index at runtime.",
            "Use lightweight local libraries such as FastAPI and scikit-learn.",
            "Do not implement web deployment, authentication, persistence, or a frontend for this assessment.",
        ],
    )

    add_h1(doc, "3. Required API Endpoints")
    add_h2(doc, "3.1 POST /index")
    doc.add_paragraph("Purpose: read all .txt files from the local docs/ folder and build the retrieval index.")
    add_bullets(
        doc,
        [
            "Read every .txt file from docs/.",
            "Split documents into smaller chunks suitable for retrieval.",
            "Build and store an in-memory retrieval index.",
            "Return basic indexing stats, including document count, chunk count, and indexed source names.",
        ],
    )
    add_table(
        doc,
        ["Field", "Type", "Description"],
        [
            ["documents_indexed", "integer", "Number of non-empty .txt files indexed"],
            ["chunks_indexed", "integer", "Number of text chunks created"],
            ["sources", "array[string]", "File names included in the index"],
        ],
        [2400, 2100, 4860],
    )

    add_h2(doc, "3.2 POST /ask")
    doc.add_paragraph("Purpose: answer a submitted question using only the indexed local documents.")
    add_bullets(
        doc,
        [
            'Accept input in the form { "question": "..." }.',
            "Retrieve the most relevant chunks from the existing in-memory index.",
            "Produce a concise extractive answer based only on retrieved content.",
            "Return the final answer and the top retrieved evidence chunks or source references.",
        ],
    )
    add_table(
        doc,
        ["Field", "Type", "Description"],
        [
            ["answer", "string", "Final answer grounded in retrieved document content"],
            ["chunks", "array[object]", "Top evidence chunks with source, chunk id, score, and text, or source references"],
        ],
        [2400, 2100, 4860],
    )

    add_h1(doc, "4. Functional Requirements")
    add_numbers(
        doc,
        [
            "The service must read all .txt files from docs/ when POST /index is called.",
            "The service must split documents into readable retrieval chunks.",
            "The service must build an in-memory retrieval index after indexing.",
            "The service must retrieve the most relevant chunks for a submitted question.",
            "The service must generate answers using only retrieved document content.",
            "The service must return evidence chunks or source references with successful answers.",
            "The service must not make up answers when evidence is weak or missing.",
            "The service must handle common error cases clearly.",
        ],
    )

    add_h1(doc, "5. Important Constraints")
    add_bullets(
        doc,
        [
            "Do not call external LLM APIs, including OpenAI, Anthropic, DeepSeek, or similar services.",
            "Do not rely on external paid services.",
            "The answer generation method should be local and extractive.",
            "If no strong evidence is found, return a clear insufficient-evidence message.",
            "Keep the implementation simple, readable, and explainable in an interview.",
        ],
    )

    add_h1(doc, "6. Required Error Handling")
    add_table(
        doc,
        ["Case", "Expected behavior"],
        [
            ["POST /ask before POST /index", 'Return a clear error such as "no index found".'],
            ["Empty question", 'Return a clear error such as "empty question".'],
            ["Missing docs/ folder", 'Return a clear error such as "docs folder not found".'],
            ["docs/ exists but has no .txt files", 'Return a clear error such as "no text documents found in docs folder".'],
            ["All document files are empty", 'Return a clear error such as "documents are empty; no index can be built".'],
            ["Weak or missing evidence", 'Return an answer such as "I cannot find enough information in the documents to answer this question."'],
        ],
        [3700, 5660],
    )

    add_h1(doc, "7. Recommended Technical Design")
    add_bullets(
        doc,
        [
            "Framework: FastAPI, because it is lightweight and provides automatic API docs.",
            "Retrieval: TF-IDF with cosine similarity via scikit-learn.",
            "Index storage: in-memory only; persistence is not required.",
            "Answer generation: extract the most relevant sentence or chunk from retrieved evidence.",
            "Evidence threshold: use a simple minimum similarity score to decide whether the evidence is strong enough.",
        ],
    )

    add_h1(doc, "8. Expected Project Structure")
    paragraph = doc.add_paragraph()
    paragraph.style = styles["Normal"]
    run = paragraph.add_run(
        "rag-api-service/\n"
        "  app.py\n"
        "  requirements.txt\n"
        "  README.md\n"
        "  docs/\n"
        "    product.txt\n"
        "    refund_policy.txt\n"
        "    shipping_policy.txt"
    )
    run.font.name = "Courier New"
    run.font.size = Pt(10)

    add_h1(doc, "9. Acceptance Criteria")
    add_bullets(
        doc,
        [
            "uvicorn app:app --reload starts the service without errors.",
            "POST /index indexes .txt files from docs/ and returns stats.",
            "POST /ask returns an answer and evidence when relevant information exists.",
            "POST /ask refuses to answer when evidence is weak or missing.",
            "POST /ask before indexing returns a clear no-index error.",
            "Empty question, missing docs folder, empty docs folder, and empty documents are handled clearly.",
            "README includes setup instructions, run command, example requests, and tradeoff notes.",
            "Code is readable and suitable for an L1 backend/RAG assessment.",
        ],
    )

    add_h1(doc, "10. README Requirements")
    add_bullets(
        doc,
        [
            "How to create and activate a virtual environment.",
            "How to install dependencies from requirements.txt.",
            "How to run the FastAPI service locally.",
            "Example curl request for POST /index.",
            "Example curl request for POST /ask.",
            "Brief notes explaining that the project uses local retrieval and extractive answering only.",
        ],
    )

    add_h1(doc, "11. Tradeoffs and Future Improvements")
    add_bullets(
        doc,
        [
            "TF-IDF is simple and local, but it may miss semantic matches with different wording.",
            "Extractive answering is safer than generative answering, but responses may sound less natural.",
            "The index is in memory, so it must be rebuilt after server restart.",
            "With more time, add tests, better chunking, persistent storage, sentence-level ranking, or local embedding models.",
        ],
    )

    doc.save(OUT_PATH)


if __name__ == "__main__":
    build_doc()
    print(OUT_PATH.resolve())
